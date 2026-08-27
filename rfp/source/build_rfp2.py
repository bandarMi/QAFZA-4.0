# -*- coding: utf-8 -*-
"""Build the Misk Leadership Competency Model and Assessment RFP from the
Professional Services RFP Template (Misk Foundation v4).

The template's styles, cover, header/footer, table of contents, numbering,
tables and standard clauses are preserved; only the project-specific content
is filled in or replaced.
"""
import copy, os

from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

import content as C

TEMPLATE = 'tpl2/template.docx'
OUT = 'out/Misk_Leadership_Competency_Model_and_Assessment_RFP_DRAFT_v2.0.docx'

BULLET_NUM = 900   # single-level bullet list  (abstract 12)
TRACK_NUM = 901    # two-level "1- / a." list  (abstract 48)

# ----------------------------------------------------------------------- utils
def set_text(p, text):
    """Replace a paragraph's text, keeping the first run's character formatting."""
    runs = p.findall(qn('w:r'))
    if not runs:
        pPr = p.find(qn('w:pPr'))
        rpr = pPr.find(qn('w:rPr')) if pPr is not None else None
        r = OxmlElement('w:r')
        if rpr is not None:
            r.append(copy.deepcopy(rpr))
        t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text
        r.append(t); p.append(r)
        return p
    first = runs[0]
    for ch in list(first):
        if ch.tag != qn('w:rPr'):
            first.remove(ch)
    t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text
    first.append(t)
    for r in runs[1:]:
        p.remove(r)
    return p

def strip_noise(el):
    for tag in ('w:bookmarkStart', 'w:bookmarkEnd', 'w:proofErr',
                'w:commentRangeStart', 'w:commentRangeEnd'):
        for e in el.findall(qn(tag)):
            el.remove(e)
    return el

def clone(ref, text=None):
    el = strip_noise(copy.deepcopy(ref))
    if text is not None:
        set_text(el, text)
    return el

def get_pPr(p):
    pPr = p.find(qn('w:pPr'))
    if pPr is None:
        pPr = OxmlElement('w:pPr'); p.insert(0, pPr)
    return pPr

def set_numbering(p, num_id, ilvl):
    pPr = get_pPr(p)
    old = pPr.find(qn('w:numPr'))
    if old is not None:
        pPr.remove(old)
    numPr = parse_xml('<w:numPr %s><w:ilvl w:val="%d"/><w:numId w:val="%d"/></w:numPr>'
                      % (nsdecls('w'), ilvl, num_id))
    ind = pPr.find(qn('w:ind'))
    (ind if ind is not None else pPr).addprevious(numPr) if ind is not None else pPr.append(numPr)
    return p

def set_indent(p, left, hanging):
    pPr = get_pPr(p)
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind'); pPr.append(ind)
    ind.set(qn('w:left'), str(left)); ind.set(qn('w:hanging'), str(hanging))
    return p

def insert_after(anchor, elements):
    for el in elements:
        anchor.addnext(el)
        anchor = el
    return anchor

def drop(el):
    el.getparent().remove(el)

# ------------------------------------------------------------------------ load
doc = Document(TEMPLATE)
body = doc.element.body
kids = [c for c in body.iterchildren() if c.tag in (qn('w:p'), qn('w:tbl'))]

def check(i, needle):
    txt = ''.join(t.text or '' for t in kids[i].iter(qn('w:t')))
    assert needle in txt, 'index %d: expected %r, got %r' % (i, needle, txt[:80])

check(11, '[Project Name]')
check(46, 'Misk Foundation (hereinafter')
check(54, 'Project Goals and Objectives')
check(56, 'Project aims at')
check(65, 'Project Duration')
check(66, 'The duration of project implementation')
check(70, 'Misk will share the candidates')
check(71, 'Misk will share the candidates')
check(75, 'Team')
check(81, 'Technical')
check(99, 'Past relevant experiences')
check(123, 'Client engagement')
check(145, 'list of suggested team members')
check(153, 'General Provisions for Project Management')
check(211, 'Additional information gets included here')

# reference elements for new content
REF_BODY   = copy.deepcopy(kids[46])    # نص عربي body paragraph
REF_EMPTY  = copy.deepcopy(kids[47])    # empty نص عربي
REF_H1     = copy.deepcopy(kids[54])    # Heading 1
REF_SUB    = copy.deepcopy(kids[153])   # نص عربي, bold + underlined sub-heading
REF_LIST   = copy.deepcopy(kids[99])    # List Paragraph carrying a numPr
REF_EVAL   = copy.deepcopy(kids[123])   # evaluation criterion (numId 45)
REF_TEAM   = copy.deepcopy(kids[76])    # Elements of the proposal - Team (numId 47)
REF_TECH   = copy.deepcopy(kids[82])    # Elements of the proposal - Technical (numId 48)
REF_OBLIG  = copy.deepcopy(kids[145])   # Key Obligations item (numId 39)

def para(text):  return clone(REF_BODY, text)
def empty():     return clone(REF_EMPTY)
def h1(text):    return clone(REF_H1, text)
def sub(text):   return clone(REF_SUB, text)

def bullet(text):
    p = clone(REF_LIST, text)
    set_numbering(p, BULLET_NUM, 0)
    set_indent(p, 720, 360)
    return p

def set_bold(r, on):
    rPr = r.find(qn('w:rPr'))
    if rPr is None:
        rPr = OxmlElement('w:rPr'); r.insert(0, rPr)
    for tag in ('w:b', 'w:bCs'):
        e = rPr.find(qn(tag))
        if e is None:
            e = OxmlElement(tag); rPr.append(e)
        e.set(qn('w:val'), 'true' if on else 'false')

def term_bullet(text, sep=' \u2013 '):
    """Bullet whose lead-in term is bold and whose explanation is regular weight."""
    head, tail = (text.split(sep, 1) if sep in text else (text, ''))
    p = bullet(head + sep if tail else head)
    r0 = p.findall(qn('w:r'))[0]
    set_bold(r0, True)
    if tail:
        r1 = copy.deepcopy(r0)
        for ch in list(r1):
            if ch.tag != qn('w:rPr'):
                r1.remove(ch)
        t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = tail
        r1.append(t)
        set_bold(r1, False)
        r0.addnext(r1)
    return p

def body_block(texts):
    """Body paragraphs separated by a blank line, as the template's own prose is."""
    out = []
    for i, t in enumerate(texts):
        if i:
            out.append(empty())
        out.append(para(t))
    return out

def track_item(lvl, text):
    p = clone(REF_LIST, text)
    set_numbering(p, TRACK_NUM, lvl)
    set_indent(p, 720 + 720 * lvl, 360)
    return p

# --------------------------------------------- register the two new list styles
numbering = doc.part.numbering_part.element

def own_abstract(source_id, new_id):
    """Copy an abstract numbering definition so the new list keeps its own counter."""
    src = None
    for el in numbering.findall(qn('w:abstractNum')):
        if el.get(qn('w:abstractNumId')) == str(source_id):
            src = el; break
    assert src is not None, 'abstractNum %s not found' % source_id
    cp = copy.deepcopy(src)
    cp.set(qn('w:abstractNumId'), str(new_id))
    for tag in ('w:nsid', 'w:tmpl', 'w:styleLink', 'w:numStyleLink'):
        for e in cp.findall(qn(tag)):
            cp.remove(e)
    src.addnext(cp)
    return new_id

for num_id, source_id, new_abstract in ((BULLET_NUM, 12, 990), (TRACK_NUM, 48, 991)):
    own_abstract(source_id, new_abstract)
    numbering.append(parse_xml('<w:num %s w:numId="%d"><w:abstractNumId w:val="%d"/></w:num>'
                               % (nsdecls('w'), num_id, new_abstract)))

# --------------------------------------------------------- 1. cover and header
set_text(kids[11], C.PROJECT_TITLE)

hdr = doc.sections[0].header
# the template's 1" first-line indent was sized for the short "[Project Name]"
# placeholder; trim it so the real project name stays on one line
for p in hdr.paragraphs:
    pPr = p._p.find(qn('w:pPr'))
    ind = pPr.find(qn('w:ind')) if pPr is not None else None
    if ind is not None and ind.get(qn('w:firstLine')):
        ind.set(qn('w:firstLine'), '360')
for p in hdr.paragraphs:
    runs = p.runs
    for i, r in enumerate(runs):
        if r.text.strip() == 'Project Name':
            runs[i - 1].text = runs[i - 1].text.replace(' [', ' ')
            r.text = C.PROJECT_TITLE
            if i + 1 < len(runs) and runs[i + 1].text.strip() == ']':
                runs[i + 1].text = ''

# --------------------------------------------------- 2. Background (new section)
blk = [h1(C.BACKGROUND_H1), empty(),
       sub(C.TRACK_H2), para(C.TRACK_INTRO)]
blk += [track_item(lvl, txt) for lvl, txt in C.TRACK_LIST]
blk += [empty(), para(C.TRACK_CLOSING_2), empty(),
        sub(C.MODEL_H2)]
blk += body_block(C.MODEL_PARAS_1)
blk += [term_bullet(t) for _, t in C.PHILOSOPHY_BULLETS]
blk += [empty()] + body_block(C.MODEL_PARAS_2)
anchor = kids[52]
for el in blk:
    anchor.addnext(el); anchor = el

# ------------------------------------------------------------ two-column tables
def two_col_table(rows, widths):
    """Clone the Key Dates table and refill it."""
    tbl = copy.deepcopy(kids[31])
    trs = tbl.findall(qn('w:tr'))
    head_ref, data_ref = copy.deepcopy(trs[0]), copy.deepcopy(trs[1])
    for tr in trs:
        tbl.remove(tr)
    for idx, values in enumerate(rows):
        tr = copy.deepcopy(head_ref if idx == 0 else data_ref)
        for tc, val in zip(tr.findall(qn('w:tc')), values):
            ps = tc.findall(qn('w:p'))
            for p in ps[1:]:
                tc.remove(p)
            set_text(strip_noise(ps[0]), val)
            if idx:                        # prose reads better left-aligned
                pPr = get_pPr(ps[0])
                jc = pPr.find(qn('w:jc'))
                if jc is None:
                    jc = OxmlElement('w:jc'); pPr.append(jc)
                jc.set(qn('w:val'), 'left')
        tbl.append(tr)
    grid = tbl.find(qn('w:tblGrid'))
    for col, w in zip(grid.findall(qn('w:gridCol')), widths):
        col.set(qn('w:w'), str(w))
    for tr in tbl.findall(qn('w:tr')):
        for tc, w in zip(tr.findall(qn('w:tc')), widths):
            tcPr = tc.find(qn('w:tcPr'))
            tcW = tcPr.find(qn('w:tcW'))
            if tcW is None:
                tcW = OxmlElement('w:tcW'); tcPr.insert(0, tcW)
            tcW.set(qn('w:w'), str(w)); tcW.set(qn('w:type'), 'dxa')
    return tbl

anchor = insert_after(anchor, [two_col_table(C.COMPETENCY_TABLE, (2900, 6640)), empty()])
anchor = insert_after(anchor, body_block(C.MODEL_PARAS_3))
anchor = insert_after(anchor, [term_bullet(t) for _, t in C.FLEX_BULLETS])
anchor = insert_after(anchor, [empty()])
anchor = insert_after(anchor, body_block(C.MODEL_PARAS_4))
anchor = insert_after(anchor, [empty()])

# ----------------------------------------------- 3. Project Goals and Objectives
blk = body_block(C.GOALS_PARAS)
blk += [empty(), sub(C.NEED_H2), para(C.NEED_INTRO)]
blk += [term_bullet(t) for _, t in C.NEED_BULLETS]
blk += [empty(), sub(C.USES_H2), para(C.USES_INTRO)]
blk += [term_bullet(t) for _, t in C.USES_BULLETS]
blk += [empty(), sub(C.POP_H2), para(C.POP_INTRO), empty()]
insert_after(kids[56], blk)
anchor = insert_after(blk[-1], [two_col_table(C.POP_TABLE, (5900, 3640))])
drop(kids[56])   # the template's "[name of project] Project aims at [...]" line

# ----------------------------------------------------------- 4. Scope of Work
scope_tbl = kids[62]

def set_table_widths(tbl, widths):
    grid = tbl.find(qn('w:tblGrid'))
    for col, w in zip(grid.findall(qn('w:gridCol')), widths):
        col.set(qn('w:w'), str(w))
    for tr in tbl.findall(qn('w:tr')):
        for tc, w in zip(tr.findall(qn('w:tc')), widths):
            tcPr = tc.find(qn('w:tcPr'))
            tcW = tcPr.find(qn('w:tcW'))
            if tcW is None:
                tcW = OxmlElement('w:tcW'); tcPr.insert(0, tcW)
            tcW.set(qn('w:w'), str(w)); tcW.set(qn('w:type'), 'dxa')
    return tbl

def set_table_indent(tbl, twips):
    tblPr = tbl.find(qn('w:tblPr'))
    ind = tblPr.find(qn('w:tblInd'))
    if ind is None:
        ind = OxmlElement('w:tblInd'); tblPr.append(ind)
    ind.set(qn('w:w'), str(twips)); ind.set(qn('w:type'), 'dxa')

# the five-column scope table needs more room than the body text column allows,
# so it runs slightly into the page margins
set_table_widths(scope_tbl, (560, 1600, 3900, 2270, 2470))
set_table_indent(scope_tbl, -580)
trs = scope_tbl.findall(qn('w:tr'))
head_tr, row_ref = trs[0], copy.deepcopy(trs[1])
for tr in trs[1:]:
    scope_tbl.remove(tr)

_cells = row_ref.findall(qn('w:tc'))
REF_TC = [copy.deepcopy(c.findall(qn('w:p'))[0]) for c in _cells]

def clear_fill(tc, valign=None):
    tcPr = tc.find(qn('w:tcPr'))
    if tcPr is None:
        return
    for shd in tcPr.findall(qn('w:shd')):
        shd.set(qn('w:fill'), 'auto')
    if valign:
        va = tcPr.find(qn('w:vAlign'))
        if va is None:
            va = OxmlElement('w:vAlign'); tcPr.append(va)
        va.set(qn('w:val'), valign)

def cell_para(ref_p, text, kind='p', size=None):
    p = strip_noise(copy.deepcopy(ref_p))
    for r in p.findall(qn('w:r')):
        p.remove(r)
    set_text(p, text)
    pPr = get_pPr(p)
    old = pPr.find(qn('w:numPr'))
    if old is not None:
        pPr.remove(old)
    jc = pPr.find(qn('w:jc'))
    if jc is None:
        jc = OxmlElement('w:jc'); pPr.append(jc)
    jc.set(qn('w:val'), 'left')
    if size is not None:                   # narrow columns need a smaller face
        for r in p.findall(qn('w:r')):
            rPr = r.find(qn('w:rPr'))
            if rPr is None:
                rPr = OxmlElement('w:rPr'); r.insert(0, rPr)
            for tag in ('w:sz', 'w:szCs'):
                e = rPr.find(qn(tag))
                if e is None:
                    e = OxmlElement(tag); rPr.append(e)
                e.set(qn('w:val'), str(size))
    if kind == 'b':
        set_numbering(p, BULLET_NUM, 0)
        set_indent(p, 227, 227)
    else:
        ind = pPr.find(qn('w:ind'))
        if ind is not None:
            pPr.remove(ind)
    return p

def fill_cell(tc, ref_p, blocks, size=None, valign=None):
    clear_fill(tc, valign)
    for p in tc.findall(qn('w:p')):
        tc.remove(p)
    for kind, text in blocks:
        tc.append(cell_para(ref_p, text, kind, size))

anchor_tr = head_tr
for idx, (no, title, details, assumptions, deliverables) in enumerate(C.SCOPE_ROWS, start=1):
    tr = copy.deepcopy(row_ref)
    tcs = tr.findall(qn('w:tc'))
    fill_cell(tcs[0], REF_TC[0], [('p', str(idx))], size=20, valign='top')
    fill_cell(tcs[1], REF_TC[2], [('p', title)], size=20, valign='top')
    fill_cell(tcs[2], REF_TC[2], details, size=20, valign='top')
    fill_cell(tcs[3], REF_TC[2], assumptions, size=20, valign='top')
    fill_cell(tcs[4], REF_TC[2], deliverables, size=20, valign='top')
    anchor_tr.addnext(tr)
    anchor_tr = tr

# ---------------------------------------------------------- 5. Project Duration
set_text(kids[66], C.DURATION_TEXT)

phase_tbl = kids[68]
trs = phase_tbl.findall(qn('w:tr'))
head_tr, row_ref = trs[0], copy.deepcopy(trs[1])
_pcells = row_ref.findall(qn('w:tc'))
REF_PC = [copy.deepcopy(c.findall(qn('w:p'))[0]) for c in _pcells]
for tr in trs[1:]:
    phase_tbl.remove(tr)

anchor_tr = head_tr
for phase, timeframe in C.PHASE_TABLE[1:]:
    tr = copy.deepcopy(row_ref)
    tcs = tr.findall(qn('w:tc'))
    fill_cell(tcs[0], REF_PC[0], [('p', phase)])
    fill_cell(tcs[1], REF_PC[1], [('p', timeframe)])
    anchor_tr.addnext(tr)
    anchor_tr = tr

# the template's leftover footnote lines about candidate interviews
drop(kids[70]); drop(kids[71])

# ------------------------------------------------- 6. Elements of the proposal
def numbered(lvl, text, num_id):
    """Elements-of-the-proposal item; both lists share the Team list's weight."""
    p = clone(REF_TEAM, text)
    set_numbering(p, num_id, lvl)
    return p

insert_after(kids[79], [numbered(lvl, t, 47) for lvl, t in C.ELEMENTS_TEAM])
for i in (76, 77, 78, 79):
    drop(kids[i])

insert_after(kids[89], [numbered(lvl, t, 48) for lvl, t in C.ELEMENTS_TECHNICAL])
for i in range(82, 90):
    drop(kids[i])

# ------------------------------------------------------- 7. Proposal Submission
insert_after(kids[99], [clone(REF_LIST, t) for t in C.TECH_PROPOSAL_EXTRA_2])

# ------------------------------------- 8. Evaluation criteria / Key Obligations
insert_after(kids[123], [clone(REF_EVAL, t) for t in C.EVAL_EXTRA_2])
insert_after(kids[145], [clone(REF_OBLIG, t) for t in C.KEY_OBLIGATIONS_EXTRA])

# -------------------------------------------------------------- 9. Appendix 2
set_text(kids[211], C.APPENDIX2_INTRO_2)
insert_after(kids[211], [term_bullet(t, sep=': ') for t in C.APPENDIX2_ITEMS])

# ------------------------------- 10. clear the template's fill-in highlighting
for part in (body, hdr._element, doc.sections[0].footer._element):
    for hl in list(part.iter(qn('w:highlight'))):
        hl.getparent().remove(hl)
for shd in body.iter(qn('w:shd')):
    if (shd.get(qn('w:fill')) or '').upper() == 'FFFF00':
        shd.set(qn('w:fill'), 'auto')

# ------------------------------------------------- 11. force TOC update on open
settings = doc.settings.element
if settings.find(qn('w:updateFields')) is None:
    uf = OxmlElement('w:updateFields'); uf.set(qn('w:val'), 'true')
    settings.append(uf)

cp = doc.core_properties
cp.title = 'Request for Proposal (RFP) - Misk Leadership Competency Model and Assessment'
cp.subject = ('Review and development of the Misk leadership competency model across the Leadership Track, '
              'and a customized assessment against it')
cp.last_modified_by = 'Misk Foundation'

os.makedirs('out', exist_ok=True)
doc.save(OUT)
print('written:', OUT, os.path.getsize(OUT), 'bytes')
