"""Fill the Misk Event Charter template for the SLS Annual Assembly 2026."""
from copy import deepcopy

from pptx import Presentation
from pptx.oxml.ns import qn

SRC = "charter_template.pptx"
OUT = "SLS_Annual_Assembly_2026_Event_Charter.pptx"
EVENT = "SLS Annual Assembly 2026"
TITLE = f"DRAFT: {EVENT} - Event Info"


def _strip_err(p):
    for rPr in p.iter(qn("a:rPr")):
        rPr.attrib.pop("err", None)


def _seed_run(p):
    """Give an empty paragraph a run, built from the run properties it ends with."""
    end = p.find(qn("a:endParaRPr"))
    run = p.makeelement(qn("a:r"), {})
    if end is not None:
        rPr = deepcopy(end)
        rPr.tag = qn("a:rPr")
        run.append(rPr)
    t = p.makeelement(qn("a:t"), {})
    t.text = ""
    run.append(t)
    if end is not None:
        end.addprevious(run)
    else:
        p.append(run)


def set_lines(cell, lines, tmpls=(0,)):
    """Rewrite a cell as one paragraph per line.

    Line i is cloned from the cell's existing paragraph tmpls[i], falling back to
    the last entry, so a cell whose first paragraph is bold keeps that emphasis.
    """
    txBody = cell.text_frame._txBody
    paras = txBody.findall(qn("a:p"))
    models = []
    for idx in tmpls:
        model = deepcopy(paras[idx])
        if not model.findall(qn("a:r")):
            _seed_run(model)
        models.append(model)
    for p in paras:
        txBody.remove(p)
    for i, line in enumerate(lines):
        p = deepcopy(models[min(i, len(models) - 1)])
        for br in p.findall(qn("a:br")):
            p.remove(br)
        runs = p.findall(qn("a:r"))
        for extra in runs[1:]:
            p.remove(extra)
        runs[0].find(qn("a:t")).text = line
        _strip_err(p)
        txBody.append(p)


def set_text(cell, text):
    set_lines(cell, [text])


def set_title(shape, text):
    """Replace the first run of a title, dropping the rest of that paragraph's runs."""
    p = shape.text_frame.paragraphs[0]._p
    runs = p.findall(qn("a:r"))
    for extra in runs[1:]:
        p.remove(extra)
    runs[0].find(qn("a:t")).text = text
    _strip_err(p)


def by_id(slide):
    return {sh.shape_id: sh for sh in slide.shapes}


def drop(shape):
    shape._element.getparent().remove(shape._element)


def set_run_color(cell, hexcolor):
    """Recolour a cell's runs. The template ships the budget amounts as white text on a
    white fill, so anything written there is invisible until this is applied."""
    for rPr in cell._tc.iter(qn("a:rPr")):
        for fill in rPr.findall(qn("a:solidFill")):
            rPr.remove(fill)
        fill = rPr.makeelement(qn("a:solidFill"), {})
        clr = rPr.makeelement(qn("a:srgbClr"), {"val": hexcolor})
        fill.append(clr)
        rPr.insert(0, fill)


def set_fill(shape, hexcolor):
    shape._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr")).set("val", hexcolor)


prs = Presentation(SRC)
s1, s2 = prs.slides[0], prs.slides[1]
b1, b2 = by_id(s1), by_id(s2)

# ------------------------------------------------------------------ slide 1
set_title(b1[19], f"{TITLE} (1/2)")

main = b1[5].table
set_text(main.cell(1, 1), EVENT)
set_text(main.cell(1, 3), "[TBC - Leadership Track]")
set_text(main.cell(2, 1), "Leadership")
set_text(main.cell(2, 3), "SLS / Leadership Track")
set_text(main.cell(3, 1), "SLS")
set_text(main.cell(3, 3), "[TBC - EMU]")
set_text(main.cell(4, 1), "Yes")
set_text(main.cell(4, 3), "City Hub (proposed)")
set_text(main.cell(5, 1), "Fellowship & 2030 Leaders")
set_text(main.cell(5, 3), "300 (venue capacity)")
set_text(main.cell(6, 1), "Private - invitation only")
set_text(main.cell(6, 3), "300 SLS members")
set_text(main.cell(7, 3), "Yes")

set_lines(b1[12].table.cell(1, 0), [
    "One-evening annual assembly for the full SLS membership. No graduation segment in 2026.",
    "Members convene at City Hub for four hours: the year's achievements shown and told by the members "
    "who delivered them, recognition of the partners, Council members, Chapter teams and contributors "
    "behind them, and the direction for 2027. Venue capacity is 300.",
], tmpls=(0, 1))

set_lines(b1[2].table.cell(1, 0), [
    "Close the SLS year by showing what the society achieved and naming the people who achieved it.",
    "Members should leave able to name three things SLS delivered in 2026, having seen their own "
    "contribution acknowledged, and knowing what 2027 asks of them. Benchmarks to beat: satisfaction "
    "of 85% or better (2024: 77%) and 85% of confirmed guests attending (2024: 320 of 468).",
], tmpls=(0, 1))

set_lines(b1[3].table.cell(1, 0), [
    "Venue: City Hub - site survey and layout for 300 across stage, seating and networking areas",
    "Agency: full pre-production, on-site delivery, dismantle and closeout",
    "Stage and set: stage and seating for 300, Year in Review gallery at the entrance",
    "Production: LED, sound, lighting, screen content, show-calling, standby power",
    "Content: the 2026 year compiled by the SLS team, three member stories, scripts, briefings",
    "Recognition: Wall of the Year, shortlist and citations prepared by the SLS team, letters",
    "Program: run of show, Isha inside the arrival window, one floor break mid-evening",
    "Registration: managed invitations, digital badge, RSVP cap and waitlist, team-confirmed count",
    "Guest experience: VIP protocol, meet and greet, seating, wayfinding, valet, parking, access",
    "F&B: dinner scaled to reconfirmed numbers, staged release",
    "Support: translation, photography, videography, cleaning, H&S, medical and security",
    "Post-event: attendee survey and post-event report",
], tmpls=(1,))

budget = b1[14].table
for row in (2, 3, 4):
    set_text(budget.cell(row, 1), "[Cost centre TBC]")
    set_text(budget.cell(row, 2), "SAR TBC")
    set_run_color(budget.cell(row, 2), "737373")
set_text(budget.cell(2, 3), "SAR TBC")

timeline = b1[16].table
set_text(timeline.cell(1, 1), "20/09/2026")
set_lines(timeline.cell(1, 3), ["12/12/2026", "6:30 - 10:30 PM"])
set_text(timeline.cell(1, 5), "14/01/2027")
set_text(timeline.cell(2, 1), "09/12/2026")
set_text(timeline.cell(2, 4), "13/12/2026")

# ------------------------------------------------------------------ slide 2
set_title(b2[25], f"{TITLE} (2/2)")

proj = b2[2].table
set_text(proj.cell(2, 2), "Approved event charter and PO issued")
set_text(proj.cell(2, 3), "20/09/2026")
set_text(proj.cell(2, 4), "15/10/2026")
set_text(proj.cell(3, 2), "Venue contract, agency award, run of show, 2026 year compiled, citations")
set_text(proj.cell(3, 3), "18/10/2026")
set_text(proj.cell(3, 4), "19/11/2026")
set_lines(proj.cell(4, 2), [
    "Invitations, registration and digital badges",
    "Build, stage, gallery and rehearsal",
    "Event day delivery and dismantle",
])
set_lines(proj.cell(4, 3), ["22/11/2026", "09/12/2026", "12/12/2026"])
set_lines(proj.cell(4, 4), ["10/12/2026", "11/12/2026", "13/12/2026"])
set_text(proj.cell(5, 2), "Post-event report, survey results and vendor closeout")
set_text(proj.cell(5, 3), "14/12/2026")
set_text(proj.cell(5, 4), "14/01/2027")

stake = b2[27].table
set_lines(stake.cell(3, 0), [
    "Leadership Track",
    "Events (EMU)",
    "Procurement/Legal",
])
set_lines(stake.cell(3, 1), [
    "[TBC] - Sponsor",
    "[TBC] - Manager",
    "[TBC] - Contracts",
])
set_lines(stake.cell(3, 2), [
    "City Hub",
    "Event agency (RFP)",
])
set_lines(stake.cell(3, 3), [
    "[TBC] - Venue ops",
    "[TBC] - Delivery",
])

risks = b2[3].table
set_text(risks.cell(2, 1), "300 seats against a full-membership invitation.")
set_text(risks.cell(2, 3), "Phased invitations with a waitlist; seating and dinner to a team-confirmed count.")
set_text(risks.cell(2, 4), "Track / EMU")
set_text(risks.cell(3, 1), "The 2026 year is not compiled in time to build the content.")
set_text(risks.cell(3, 3), "SLS team compiles from its own records by 19/11; Chapters asked only for gaps.")
set_text(risks.cell(3, 4), "Leadership Track")

# Severity: both entries are risks, so replace the second row's issue icon with a
# risk icon, and set the dots to High / Medium.
set_fill(b2[4], "C00000")   # row 1 severity - High
set_fill(b2[5], "FFB500")   # row 2 severity - Medium

risk_icon = deepcopy(b2[6]._element)
src_y = int(b2[6]._element.spPr.find(qn("a:xfrm")).find(qn("a:off")).get("y"))
delta = int(b2[5]._element.spPr.find(qn("a:xfrm")).find(qn("a:off")).get("y")) - \
         int(b2[4]._element.spPr.find(qn("a:xfrm")).find(qn("a:off")).get("y"))
risk_icon.spPr.find(qn("a:xfrm")).find(qn("a:off")).set("y", str(src_y + delta))
risk_icon.nvSpPr.cNvPr.set("id", "300")
risk_icon.nvSpPr.cNvPr.set("name", "Risk icon row 2")
drop(b2[7])                                     # the "!" issue icon
b2[6]._element.getparent().append(risk_icon)

# Level of attendees: keep Key CEOs, Partners, Beneficiaries.
drop(b2[33])    # Public
# CEO / DCEO required: CEO only.
drop(b2[10])    # DCEO

prs.save(OUT)
print("wrote", OUT)
